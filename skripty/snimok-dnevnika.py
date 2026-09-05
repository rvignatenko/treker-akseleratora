#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Снимок официального «Дневника команды» и контроль того, ведёт ли его команда.

Зачем. Дневник команды — документ, который по правилам ведёт сама команда,
а на практике почти всегда ведёт трекер. Понять, включилась ли команда,
на глаз нельзя: файл лежит на общем диске и меняется молча. Скрипт снимает
состояние всех листов, отличает записи трекера от записей команды по цвету
шрифта, сравнивает с прошлым снимком и дописывает разницу в реестр.

Только читает облачный файл. Ничего в нём не меняет и наружу не выкладывает.

    python3 snimok-dnevnika.py "<команда>"
    python3 snimok-dnevnika.py --vse

Требуется: pip install openpyxl

Автор: Роман Игнатенко, https://t.me/ignatenko_roman
"""
import json, os, re, sys, time, unicodedata as ud
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from nastrojka import prochitat, papka_komandy, najti_komandu  # noqa: E402

try:
    import openpyxl
except ImportError:
    raise SystemExit('Нужен openpyxl: pip3 install openpyxl')

NFC = lambda s: ud.normalize('NFC', str(s))
KOLONKA = re.compile(r'^([A-Z]+)(\d+)$')

NASH, KOMANDA, SMESH, NEYASNO = 'nash', 'komanda', 'smeshannoe', 'neyasno'
ZNACHKI = {NASH: '🔴', KOMANDA: '⚫️', SMESH: '🟠', NEYASNO: '⚪️'}
PODPISI = {NASH: 'наше (трекер)', KOMANDA: 'команда или трекер на встрече',
           SMESH: 'смешанное', NEYASNO: 'третий цвет — посмотреть глазами'}


CHERNYE = {'FF000000', '00000000', '000000', 'FF1F1F1F'}


def cvet_runa(shrift, cvet_trekera: str):
    """Трекер помечает свои дописи одним явным цветом. Всё остальное — обычный
    текст, то есть запись команды или ваша на встрече.

    Отдельно помечаем «неясно» ровно один случай: цвет задан явно, но он и не ваш,
    и не чёрный. Это чей-то третий цвет, и его стоит посмотреть глазами, а не
    записывать молча в чью-то пользу."""
    if shrift is None or shrift.color is None:
        return KOMANDA
    c = shrift.color
    if getattr(c, 'type', None) != 'rgb' or not c.rgb:
        return KOMANDA           # цвет из темы оформления — это обычный текст
    rgb = str(c.rgb).upper()
    if rgb == cvet_trekera.upper():
        return NASH
    return KOMANDA if rgb in CHERNYE else NEYASNO


def avtorstvo(cell, cvet_trekera: str):
    """У ячейки с rich text куски могут быть разного цвета."""
    znachenie = getattr(cell, 'value', None)
    bloki = getattr(znachenie, 'rich', None)
    if bloki:
        cveta = {cvet_runa(getattr(b, 'font', None), cvet_trekera) for b in bloki}
        cveta.discard(NEYASNO)
        if len(cveta) > 1:
            return SMESH
        if cveta:
            return next(iter(cveta))
        return NEYASNO
    return cvet_runa(cell.font, cvet_trekera)


def tekst(cell) -> str:
    v = cell.value
    if v is None:
        return ''
    if hasattr(v, 'rich') and v.rich:
        return ''.join(getattr(b, 'text', '') for b in v.rich).strip()
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    return str(v).strip()


def snyat(put_k_fajlu: Path, cvet_trekera: str) -> dict:
    kniga = openpyxl.load_workbook(put_k_fajlu, rich_text=True, data_only=False)
    listy = {}
    for name in kniga.sheetnames:
        list_ = kniga[name]
        yachejki = {}
        for ryad in list_.iter_rows():
            for cell in ryad:
                t = tekst(cell)
                if t:
                    yachejki[cell.coordinate] = [t, avtorstvo(cell, cvet_trekera)]
        if yachejki:
            listy[NFC(name)] = yachejki
    kniga.close()
    return listy


def najti_dnevnik(papka: Path, kusok_imeni: str) -> Path:
    kandidaty = sorted(
        (f for f in papka.iterdir()
         if f.suffix.lower() in ('.xlsx', '.xlsm') and not f.name.startswith('~$')
         and NFC(kusok_imeni).lower() in NFC(f.name).lower()),
        key=lambda f: f.stat().st_mtime, reverse=True)
    if not kandidaty:
        raise SystemExit(f'В папке команды на диске нет файла «{kusok_imeni}*.xlsx»: {papka}\n'
                         'Либо команда его ещё не завела, либо облачный диск не смонтирован.')
    if len(kandidaty) > 1:
        print(f'   (в папке несколько дневников, беру самый свежий: {kandidaty[0].name})')
    return kandidaty[0]


def raznica(bylo: dict, stalo: dict) -> dict:
    dobavleno, izmeneno, udaleno = [], [], []
    for list_, yachejki in stalo.items():
        prezhde = bylo.get(list_, {})
        for koord, (t, avt) in yachejki.items():
            if koord not in prezhde:
                dobavleno.append((list_, koord, t, avt))
            elif prezhde[koord][0] != t:
                izmeneno.append((list_, koord, t, avt, prezhde[koord][0]))
    for list_, yachejki in bylo.items():
        for koord, (t, _) in yachejki.items():
            if koord not in stalo.get(list_, {}):
                udaleno.append((list_, koord, t))
    return {'dobavleno': dobavleno, 'izmeneno': izmeneno, 'udaleno': udaleno}


def lipkoe_avtorstvo(bylo: dict, stalo: dict) -> None:
    """Если текст ячейки не менялся, авторство наследуем от прошлого снимка:
    цвет мог быть переопределён темой оформления, а автор от этого не изменился."""
    for list_, yachejki in stalo.items():
        prezhde = bylo.get(list_, {})
        for koord, para in yachejki.items():
            staroe = prezhde.get(koord)
            if staroe and staroe[0] == para[0] and para[1] in (KOMANDA, NEYASNO):
                para[1] = staroe[1]


def svodka(stalo: dict, cfg_dnevnik: dict) -> list:
    stroki = []
    for list_, yachejki in stalo.items():
        schet = {}
        for _, avt in yachejki.values():
            schet[avt] = schet.get(avt, 0) + 1
        chasti = ' '.join(f'{ZNACHKI[a]}{n}' for a, n in sorted(schet.items()) if n)
        stroki.append(f'- **{list_}** — заполнено ячеек {len(yachejki)}: {chasti}')
        kolonki = cfg_dnevnik.get('kolonki_vstrech') or {}
        if kolonki and any(NFC(str(v)).lower() in list_.lower() for v in [cfg_dnevnik.get('list_vstrechi', '')]):
            po_kolonkam = {}
            for koord, (_, avt) in yachejki.items():
                m = KOLONKA.match(koord)
                if m:
                    po_kolonkam.setdefault(m.group(1), []).append(avt)
            hvost = []
            for bukva, podpis in kolonki.items():
                if bukva in po_kolonkam:
                    a = po_kolonkam[bukva]
                    hvost.append(f'{podpis}: ' + ' '.join(
                        f'{ZNACHKI[k]}{a.count(k)}' for k in (KOMANDA, NASH, SMESH, NEYASNO) if a.count(k)))
            if hvost:
                stroki.append('  - ' + '; '.join(hvost))
    return stroki


def obrabotat(cfg: dict, komanda: dict) -> int:
    d = cfg.get('dnevnik_komandy', {})
    cvet = d.get('cvet_trekera', 'FFFF0000')
    setevaya = komanda.get('papka_setevaya')
    if not setevaya:
        print(f'{komanda["nazvanie"]}: не указана сетевая папка — пропускаю.')
        return 2
    setevaya = Path(setevaya)
    if not setevaya.exists():
        print(f'{komanda["nazvanie"]}: папка на диске недоступна ({setevaya}). '
              'Облачный диск не смонтирован?')
        return 2

    dnevnik = najti_dnevnik(setevaya, d.get('fajl_soderzhit', 'Дневник команды'))
    imena = cfg.get('imena', {})
    kontrol = (papka_komandy(cfg, komanda) / cfg['struktura_komandy']['dnevniki']
               / imena.get('papka_kontrolya', '_kontrol-dnevnika'))
    snimki = kontrol / imena.get('papka_snimkov', 'snimki')
    snimki.mkdir(parents=True, exist_ok=True)
    reestr = kontrol / cfg.get('imena', {}).get('reestr_dnevnika', 'РЕЕСТР.md')

    zamok = kontrol / '.zamok'
    if zamok.exists() and time.time() - zamok.stat().st_mtime < 300:
        print(f'{komanda["nazvanie"]}: рядом уже идёт снимок — подожди минуту.')
        return 4
    zamok.write_text(str(os.getpid()), encoding='utf-8')
    try:
        prezhnie = sorted(snimki.glob('20*.json'))
        bylo = {}
        if prezhnie:
            with open(prezhnie[-1], encoding='utf-8') as fh:
                bylo = json.load(fh).get('listy', {})

        stalo = snyat(dnevnik, cvet)
        lipkoe_avtorstvo(bylo, stalo)
        r = raznica(bylo, stalo)
        vsego = len(r['dobavleno']) + len(r['izmeneno']) + len(r['udaleno'])

        metka = datetime.now().strftime('%Y-%m-%dT%H-%M-%S')
        with open(snimki / f'{metka}.json', 'w', encoding='utf-8') as fh:
            json.dump({'metka': metka, 'komanda': komanda['nazvanie'],
                       'fajl': dnevnik.name, 'listy': stalo}, fh, ensure_ascii=False, indent=1)

        L = [f'\n## {datetime.now():%d.%m.%Y %H:%M} — снимок',
             f'Файл: `{dnevnik.name}`, изменён {datetime.fromtimestamp(dnevnik.stat().st_mtime):%d.%m.%Y %H:%M}.']
        if not prezhnie:
            L.append('Первый снимок — сравнивать не с чем.')
        elif vsego == 0:
            L.append(f'С прошлого снимка ({prezhnie[-1].stem}) изменений нет.')
        else:
            L.append(f'**Изменения с прошлого снимка ({prezhnie[-1].stem}):**')
            for zagolovok, kljuch in (('Добавлено', 'dobavleno'), ('Изменено', 'izmeneno')):
                gruppy = {}
                for zapis in r[kljuch]:
                    gruppy.setdefault(zapis[3], []).append(zapis)
                for avt, spisok in gruppy.items():
                    L.append(f'\n{ZNACHKI[avt]} **{zagolovok} — {PODPISI[avt]}:**')
                    for zapis in spisok:
                        L.append(f'- `{zapis[0]}!{zapis[1]}` — {zapis[2][:160]}')
            if r['udaleno']:
                L.append('\n**Удалено:**')
                for list_, koord, t in r['udaleno']:
                    L.append(f'- `{list_}!{koord}` — было: {t[:120]}')
        L.append('\n**Текущее заполнение:**')
        L += svodka(stalo, d)
        L.append('\n---')

        if not reestr.exists():
            shapka = (f'# Реестр заполнения дневника — {komanda["nazvanie"]}\n\n'
                      'Внутренний контроль трекера. На общий диск НЕ выкладывается: '
                      'команда и организаторы этот файл не видят.\n\n'
                      'Обозначения: '
                      + ' · '.join(f'{ZNACHKI[k]} {PODPISI[k]}' for k in (KOMANDA, NASH, SMESH, NEYASNO))
                      + '.\n\nПоявилось новое ⚫️ — уточни у себя же: это вписал ты на встрече '
                        'или команда сама. Чёрное не от тебя означает, что команда взялась '
                        'вести дневник, и это стоит беречь.\n\n'
                        'Снимок: `python3 skripty/snimok-dnevnika.py "<команда>"`\n')
            with open(reestr, 'w', encoding='utf-8') as fh:
                fh.write(shapka)
        with open(reestr, 'a', encoding='utf-8') as fh:
            fh.write('\n'.join(L) + '\n')

        print(f'OK · {komanda["nazvanie"]} · снимок {metka} · изменений: {vsego}')
        print(f'     {reestr}')
        return 0
    finally:
        zamok.unlink(missing_ok=True)


def main() -> int:
    cfg = prochitat()
    argv = [a for a in sys.argv[1:] if not a.startswith('--')]
    kljuchi = {a for a in sys.argv[1:] if a.startswith('--')}
    if '--vse' in kljuchi or not argv:
        komandy = cfg['komandy']
    else:
        komandy = [najti_komandu(cfg, argv[0])]
    kody = [obrabotat(cfg, k) for k in komandy]
    return max(kody) if kody else 0


if __name__ == '__main__':
    sys.exit(main())
